# -*- coding: utf-8 -*-
"""
把项目报告的 Markdown 转成 Word 文档
=====================================
处理：标题层级 / 表格 / 代码块 / 图片 / 粗体 / 段落
排版：中文用宋体（正文）+ 黑体（标题），西文用 Times New Roman
"""

import sys, io, re
sys.stdout.reconfigure(encoding="utf-8")
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path(r"D:\college documents\python数据分析 项目")
SRC = BASE / "项目报告_题1无人机遥测.md"
OUT = BASE / "项目报告_题1无人机遥测.docx"

FONT_CN_BODY = "宋体"
FONT_CN_HEAD = "黑体"
FONT_EN = "Times New Roman"


def set_run_font(run, cn=FONT_CN_BODY, en=FONT_EN, size=None, bold=None, color=None):
    run.font.name = en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, style=None, size=10.5, bold=False, align=None,
             cn=FONT_CN_BODY, first_indent=True, space_after=6):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.4
    if first_indent and style is None:
        pf.first_line_indent = Pt(21)   # 首行缩进两字符

    # 解析 **bold** 片段
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for seg in parts:
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2])
            set_run_font(r, cn=cn, size=size, bold=True)
        else:
            r = p.add_run(seg)
            set_run_font(r, cn=cn, size=size, bold=bold)
    return p


def add_heading(doc, text, level):
    sizes = {1: 18, 2: 15, 3: 12.5, 4: 11.5}
    h = doc.add_heading(level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = h.paragraph_format
    pf.space_before = Pt(14 if level <= 2 else 10)
    pf.space_after = Pt(8)
    r = h.add_run(text)
    set_run_font(r, cn=FONT_CN_HEAD, en=FONT_EN, size=sizes.get(level, 11.5), bold=True,
                 color=RGBColor(0, 0, 0))
    return h


def add_code_block(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.1
        pf.left_indent = Cm(0.6)
        r = p.add_run(ln if ln else " ")
        set_run_font(r, cn="Consolas", en="Consolas", size=9)
        r.font.color.rgb = RGBColor(0x1F, 0x2A, 0x33)
    # 代码块后加一点空隙
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def set_cell_text(cell, text, bold=False, size=9, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    inner = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    r = p.add_run(inner)
    set_run_font(r, cn=FONT_CN_BODY, size=size, bold=bold)


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            val = row[ci] if ci < len(row) else ""
            # 表头加粗 + 浅灰底
            set_cell_text(cells[ci], val, bold=(ri == 0))
            if ri == 0:
                tcPr = cells[ci]._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "F2F4F7")
                tcPr.append(shd)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_image(doc, path, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run()
    try:
        run.add_picture(str(path), width=Cm(width_cm))
    except Exception as e:
        r = p.add_run(f"[图片缺失: {path.name}]")
        set_run_font(r, size=9, color=RGBColor(0x99, 0x00, 0x00))
        print(f"  ! 图片插入失败 {path.name}: {e}")


def parse_table(lines, i):
    """从第 i 行开始收集表格，返回 (rows, next_index)"""
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip()
        # 跳过分隔行 |---|---|
        if re.match(r"^\|[\s:\-\|]+\|$", raw):
            i += 1
            continue
        cells = [c.strip() for c in raw.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def build():
    md = SRC.read_text(encoding="utf-8")
    lines = md.split("\n")

    doc = Document()

    # ---- 页面设置 ----
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.5)

    # 默认样式
    st = doc.styles["Normal"]
    st.font.name = FONT_EN
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN_BODY)

    i = 0
    in_code = False
    code_buf = []
    cover_done = False

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        # 代码块
        if s.startswith("```"):
            if in_code:
                add_code_block(doc, code_buf)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(ln)
            i += 1
            continue

        # 空行
        if not s:
            i += 1
            continue

        # 分隔线 —— 在封面后插入分页
        if re.match(r"^-{3,}$", s):
            if not cover_done:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                cover_done = True
            i += 1
            continue

        # 图片
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)$", s)
        if m:
            rel = m.group(1)
            add_image(doc, BASE / rel)
            i += 1
            continue

        # 表格
        if s.startswith("|"):
            rows, ni = parse_table(lines, i)
            add_table(doc, rows)
            i = ni
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if not cover_done:
                # 封面：一级标题居中放大
                if level == 1:
                    p = add_para(doc, text, size=20, bold=True,
                                 align=WD_ALIGN_PARAGRAPH.CENTER, cn=FONT_CN_HEAD,
                                 first_indent=False, space_after=24)
                    i += 1
                    continue
            add_heading(doc, text, level)
            i += 1
            continue

        # 封面信息行（**报告题目：** xxx）
        if not cover_done and s.startswith("**") and "：" in s:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            parts = re.split(r"(\*\*[^*]+\*\*)", s)
            for seg in parts:
                if not seg:
                    continue
                if seg.startswith("**") and seg.endswith("**"):
                    r = p.add_run(seg[2:-2])
                    set_run_font(r, size=12, bold=True)
                else:
                    r = p.add_run(seg)
                    set_run_font(r, size=12)
            i += 1
            continue

        # 普通段落
        add_para(doc, s)
        i += 1

    if in_code and code_buf:
        add_code_block(doc, code_buf)

    doc.save(str(OUT))
    print(f"已生成: {OUT}")
    print(f"大小: {OUT.stat().st_size/1024:.1f} KB")


if __name__ == "__main__":
    build()
